"""End-to-end audio -> transcript -> refined text -> Word document pipeline.

Takes a single audio file or a zip archive of several, and produces one
reviewed Word document per recording.

Stages:
  1. extract    unpack the input (zip -> many files, or a single file as-is)
                into tmp/extracted
  2. normalize  convert every audio file to 16 kHz mono WAV via ffmpeg
  3. transcribe run Whisper large-v3-turbo (MLX) over each file
  4. refine     clean each transcript with a local Ollama SLM
  5. docx       write one Word document per audio file

Each stage caches its result in output/transcripts.json, so re-running skips
work that is already done. Use --force to redo everything.

Usage:
    uv run python pipeline.py --input recordings.zip
    uv run python pipeline.py --input one_interview.m4a
    uv run python pipeline.py                          # auto-picks the lone .zip in the project root, if any
    uv run python pipeline.py --input recordings.zip --force
    uv run python pipeline.py --only transcribe refine
    uv run python pipeline.py --input recordings.zip --skip-refine
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
import time
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import config

STAGES = ("extract", "normalize", "transcribe", "refine", "docx")


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------

def log(stage: str, msg: str) -> None:
    print(f"[{stage:<9}] {msg}", flush=True)


def require(binary: str) -> None:
    if shutil.which(binary) is None:
        sys.exit(f"Required binary '{binary}' not found on PATH.")


def human_duration(seconds: float) -> str:
    m, s = divmod(int(round(seconds)), 60)
    return f"{m:d}:{s:02d}"


def load_state() -> dict:
    if config.TRANSCRIPTS_JSON.exists():
        return json.loads(config.TRANSCRIPTS_JSON.read_text(encoding="utf-8"))
    return {"generated_at": None, "asr_model": None, "llm_model": None, "items": {}}


def save_state(state: dict) -> None:
    config.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    state["generated_at"] = datetime.now().isoformat(timespec="seconds")
    state["asr_model"] = config.ASR_MODEL
    state["llm_model"] = config.LLM_MODEL
    config.TRANSCRIPTS_JSON.write_text(
        json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8"
    )


# ---------------------------------------------------------------------------
# stage 1 - extract
# ---------------------------------------------------------------------------

def resolve_input(explicit: Path | None) -> Path:
    """Figure out what to process: an explicit path, or the lone zip in ROOT."""
    if explicit is not None:
        if not explicit.exists():
            sys.exit(f"Input not found: {explicit}")
        return explicit

    zips = sorted(config.ROOT.glob("*.zip"))
    if len(zips) == 1:
        return zips[0]
    if len(zips) > 1:
        sys.exit(
            f"Multiple zip files found in {config.ROOT}; "
            f"pick one with --input <path>."
        )
    sys.exit(
        "No input given and no .zip file found in the project root. "
        "Pass a zip archive of recordings or a single audio file, e.g.\n"
        "  uv run python pipeline.py --input recordings.zip\n"
        "  uv run python pipeline.py --input interview.m4a"
    )


def stage_extract(input_path: Path, force: bool) -> list[Path]:
    """Unpack the input into tmp/extracted.

    Accepts either a zip archive of recordings (skipping macOS resource-fork
    junk) or a single audio file, which is just copied in as-is so every
    later stage can treat "one file" and "many files" identically.
    """
    if force and config.EXTRACT_DIR.exists():
        shutil.rmtree(config.EXTRACT_DIR)
    config.EXTRACT_DIR.mkdir(parents=True, exist_ok=True)

    if input_path.suffix.lower() != ".zip":
        if input_path.suffix.lower() not in config.AUDIO_EXTENSIONS:
            sys.exit(f"Unrecognised input type: {input_path.suffix} ({input_path})")
        target = config.EXTRACT_DIR / input_path.name
        if force or not target.exists():
            shutil.copyfile(input_path, target)
        log("extract", f"single audio file: {input_path.name}")
        return [target]

    extracted = 0
    with zipfile.ZipFile(input_path) as zf:
        for info in zf.infolist():
            name = info.filename
            if info.is_dir():
                continue
            # macOS archives carry a parallel __MACOSX/._name tree of AppleDouble
            # metadata files. They share the real file's extension but hold no audio.
            if "__MACOSX" in name or Path(name).name.startswith("._"):
                continue
            if Path(name).suffix.lower() not in config.AUDIO_EXTENSIONS:
                log("extract", f"skipping non-audio entry: {name}")
                continue
            # Flatten any directory structure, keep the original basename so a
            # transcript can always be traced back to its source recording.
            target = config.EXTRACT_DIR / Path(name).name
            if target.exists() and not force:
                continue
            with zf.open(info) as src, open(target, "wb") as dst:
                shutil.copyfileobj(src, dst)
            extracted += 1

    files = sorted(
        p for p in config.EXTRACT_DIR.iterdir()
        if p.is_file() and p.suffix.lower() in config.AUDIO_EXTENSIONS
    )
    log("extract", f"{extracted} newly extracted, {len(files)} audio files total")
    return files


# ---------------------------------------------------------------------------
# stage 2 - normalize
# ---------------------------------------------------------------------------

def probe(path: Path) -> dict:
    """Return codec/sample-rate/channels/duration for an audio file."""
    out = subprocess.run(
        ["ffprobe", "-v", "error",
         "-select_streams", "a:0",
         "-show_entries", "stream=codec_name,sample_rate,channels",
         "-show_entries", "format=duration",
         "-of", "json", str(path)],
        capture_output=True, text=True, check=True,
    ).stdout
    data = json.loads(out)
    streams = data.get("streams") or [{}]
    stream = streams[0]
    duration = data.get("format", {}).get("duration")
    return {
        "codec": stream.get("codec_name"),
        "sample_rate": int(stream["sample_rate"]) if stream.get("sample_rate") else None,
        "channels": stream.get("channels"),
        "duration": float(duration) if duration else None,
    }


def stage_normalize(files: list[Path], state: dict, force: bool) -> None:
    """Convert everything to 16 kHz mono PCM WAV, which is what Whisper wants."""
    config.NORMALIZED_DIR.mkdir(parents=True, exist_ok=True)

    for path in files:
        key = path.name
        item = state["items"].setdefault(key, {})
        wav = config.NORMALIZED_DIR / f"{path.stem}.wav"

        if wav.exists() and not force and item.get("normalized"):
            continue

        try:
            info = probe(path)
        except subprocess.CalledProcessError as exc:
            log("normalize", f"FAILED to probe {key}: {exc.stderr.strip()[:200]}")
            item["error"] = "probe failed"
            continue

        if info["codec"] is None:
            log("normalize", f"no audio stream in {key}, skipping")
            item["error"] = "no audio stream"
            continue

        already_ok = (
            path.suffix.lower() == ".wav"
            and info["sample_rate"] == config.TARGET_SAMPLE_RATE
            and info["channels"] == config.TARGET_CHANNELS
        )
        action = "copy (already 16k mono)" if already_ok else (
            f"convert {info['codec']} {info['sample_rate']}Hz "
            f"{info['channels']}ch -> 16k mono"
        )

        if already_ok:
            shutil.copyfile(path, wav)
        else:
            # -ac 1 downmix, -ar 16000 resample, signed 16-bit PCM.
            proc = subprocess.run(
                ["ffmpeg", "-v", "error", "-y", "-i", str(path),
                 "-ac", str(config.TARGET_CHANNELS),
                 "-ar", str(config.TARGET_SAMPLE_RATE),
                 "-c:a", "pcm_s16le", str(wav)],
                capture_output=True, text=True,
            )
            if proc.returncode != 0:
                log("normalize", f"FAILED {key}: {proc.stderr.strip()[:200]}")
                item["error"] = "ffmpeg conversion failed"
                continue

        item.update({
            "source_file": key,
            "source_format": info["codec"],
            "source_sample_rate": info["sample_rate"],
            "source_channels": info["channels"],
            "duration_seconds": round(info["duration"], 2) if info["duration"] else None,
            "normalized": str(wav.relative_to(config.ROOT)),
        })
        item.pop("error", None)
        log("normalize", f"{key}: {action}")


# ---------------------------------------------------------------------------
# stage 3 - transcribe
# ---------------------------------------------------------------------------

def stage_transcribe(state: dict, force: bool) -> None:
    import mlx_whisper

    pending = [
        (k, v) for k, v in state["items"].items()
        if v.get("normalized") and (force or not v.get("raw_transcript"))
    ]
    if not pending:
        log("transcribe", "nothing to do (all cached)")
        return

    log("transcribe", f"{len(pending)} file(s) with {config.ASR_MODEL}")
    for key, item in pending:
        wav = config.ROOT / item["normalized"]
        started = time.time()
        try:
            result = mlx_whisper.transcribe(
                str(wav),
                path_or_hf_repo=config.ASR_MODEL,
                language=config.ASR_LANGUAGE,
                verbose=False,
            )
        except Exception as exc:  # noqa: BLE001 - keep the batch going
            log("transcribe", f"FAILED {key}: {exc}")
            item["error"] = f"asr failed: {exc}"
            continue

        elapsed = time.time() - started
        item["raw_transcript"] = result["text"].strip()
        item["segments"] = [
            {"start": round(s["start"], 2), "end": round(s["end"], 2),
             "text": s["text"].strip()}
            for s in result.get("segments", [])
        ]
        item["asr_seconds"] = round(elapsed, 2)
        item.pop("error", None)

        dur = item.get("duration_seconds") or 0
        speed = f"{dur / elapsed:.1f}x realtime" if elapsed > 0 and dur else ""
        log("transcribe", f"{key}: {len(item['raw_transcript'])} chars "
                          f"in {elapsed:.1f}s {speed}")
        save_state(state)  # checkpoint after every file


# ---------------------------------------------------------------------------
# stage 4 - refine
# ---------------------------------------------------------------------------

def stage_refine(state: dict, force: bool) -> None:
    import ollama

    pending = [
        (k, v) for k, v in state["items"].items()
        if v.get("raw_transcript") and (force or not v.get("refined_transcript"))
    ]
    if not pending:
        log("refine", "nothing to do (all cached)")
        return

    log("refine", f"{len(pending)} file(s) with {config.LLM_MODEL}")
    for key, item in pending:
        prompt = config.REFINE_PROMPT.format(transcript=item["raw_transcript"])
        started = time.time()
        try:
            response = ollama.chat(
                model=config.LLM_MODEL,
                messages=[
                    {"role": "system", "content": config.REFINE_SYSTEM},
                    {"role": "user", "content": prompt},
                ],
                # This model reasons by default; thinking adds minutes per file
                # and nothing to the output quality for a cleanup task.
                think=False,
                options=config.LLM_OPTIONS,
            )
        except Exception as exc:  # noqa: BLE001
            log("refine", f"FAILED {key}: {exc}")
            item["refine_error"] = str(exc)
            continue

        text = response["message"]["content"].strip()
        # Strip a markdown fence if the model wrapped its answer in one.
        if text.startswith("```"):
            lines = text.splitlines()
            text = "\n".join(lines[1:-1] if lines[-1].strip() == "```" else lines[1:]).strip()

        item["refined_transcript"] = text
        item["refine_seconds"] = round(time.time() - started, 2)
        item.pop("refine_error", None)
        log("refine", f"{key}: {len(text)} chars in {item['refine_seconds']:.1f}s")
        save_state(state)


# ---------------------------------------------------------------------------
# stage 5 - docx
# ---------------------------------------------------------------------------

def stage_docx(state: dict) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    config.DOCX_DIR.mkdir(parents=True, exist_ok=True)
    written = 0

    for key, item in sorted(state["items"].items()):
        body = item.get("refined_transcript") or item.get("raw_transcript")
        if not body:
            continue

        doc = Document()
        doc.core_properties.title = f"Transcrição — {key}"
        doc.core_properties.comments = f"Ficheiro de origem: {key}"

        doc.add_heading("Transcrição de Áudio", level=0)

        # Provenance block: this is what ties the document back to its recording.
        meta = doc.add_paragraph()
        meta.add_run("Ficheiro de origem: ").bold = True
        meta.add_run(key)
        meta.add_run("\nDuração: ").bold = True
        meta.add_run(human_duration(item["duration_seconds"])
                     if item.get("duration_seconds") else "desconhecida")
        meta.add_run("\nFormato original: ").bold = True
        meta.add_run(f"{item.get('source_format', '?')} · "
                     f"{item.get('source_sample_rate', '?')} Hz · "
                     f"{item.get('source_channels', '?')} canal(is)")
        meta.add_run("\nModelo de transcrição: ").bold = True
        meta.add_run(config.ASR_MODEL)
        meta.add_run("\nModelo de revisão: ").bold = True
        meta.add_run(config.LLM_MODEL if item.get("refined_transcript")
                     else "— (texto bruto, sem revisão)")
        meta.add_run("\nGerado em: ").bold = True
        meta.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))
        for run in meta.runs:
            run.font.size = Pt(9)

        note = doc.add_paragraph()
        note_run = note.add_run(
            "Documento gerado automaticamente. Trechos marcados com [?] são "
            "incertos e requerem confirmação humana."
        )
        note_run.italic = True
        note_run.font.size = Pt(9)
        note_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        doc.add_paragraph()
        doc.add_heading("Transcrição revista" if item.get("refined_transcript")
                        else "Transcrição (bruta)", level=1)

        for block in [b.strip() for b in body.split("\n") if b.strip()]:
            para = doc.add_paragraph(block)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            para.paragraph_format.space_after = Pt(8)

        # Keep the unedited ASR output alongside it so a reviewer can check
        # any correction the model made against what was actually heard.
        if item.get("refined_transcript") and item.get("raw_transcript"):
            doc.add_page_break()
            doc.add_heading("Anexo — transcrição automática original", level=1)
            anexo = doc.add_paragraph(item["raw_transcript"])
            anexo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in anexo.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        out = config.DOCX_DIR / f"{Path(key).stem}.docx"
        doc.save(out)
        item["docx"] = str(out.relative_to(config.ROOT))
        written += 1
        log("docx", f"{out.name}")

    log("docx", f"{written} document(s) written to {config.DOCX_DIR.relative_to(config.ROOT)}")


# ---------------------------------------------------------------------------
# entrypoint
# ---------------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--input", type=Path, default=None, metavar="PATH",
                    help="a zip archive of recordings, or a single audio file. "
                         "If omitted, uses the only .zip in the project root.")
    ap.add_argument("--force", action="store_true",
                    help="redo every stage, ignoring cached results")
    ap.add_argument("--only", nargs="+", choices=STAGES, metavar="STAGE",
                    help=f"run only these stages ({', '.join(STAGES)})")
    ap.add_argument("--skip-refine", action="store_true",
                    help="transcribe only, no LLM cleanup pass")
    args = ap.parse_args()

    stages = set(args.only) if args.only else set(STAGES)
    if args.skip_refine:
        stages.discard("refine")

    require("ffmpeg")
    require("ffprobe")

    state = load_state()
    t0 = time.time()

    files: list[Path] = []
    if "extract" in stages:
        input_path = resolve_input(args.input)
        files = stage_extract(input_path, args.force)
    else:
        files = sorted(
            p for p in config.EXTRACT_DIR.glob("*")
            if p.is_file() and p.suffix.lower() in config.AUDIO_EXTENSIONS
        ) if config.EXTRACT_DIR.exists() else []

    if "normalize" in stages:
        stage_normalize(files, state, args.force)
        save_state(state)

    if "transcribe" in stages:
        stage_transcribe(state, args.force)
        save_state(state)

    if "refine" in stages:
        stage_refine(state, args.force)
        save_state(state)

    if "docx" in stages:
        stage_docx(state)
        save_state(state)

    ok = sum(1 for v in state["items"].values() if v.get("raw_transcript"))
    audio = sum(v.get("duration_seconds") or 0 for v in state["items"].values())
    log("done", f"{ok}/{len(state['items'])} transcribed · "
                f"{human_duration(audio)} of audio · "
                f"{time.time() - t0:.1f}s wall clock")
    log("done", f"JSON: {config.TRANSCRIPTS_JSON.relative_to(config.ROOT)}")


if __name__ == "__main__":
    main()
