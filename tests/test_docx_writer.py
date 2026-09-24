"""Tests for write_docx, focused on the Phase 3 summary integration.

Regression coverage for a real bug found while running Phase 3 against real
recordings: OOXML core properties (subject/keywords/title) cap out at 255
unicode characters, and an LLM-written 3-6 sentence description routinely
exceeds that -- python-docx raises ValueError past the limit.
"""

from __future__ import annotations

from docx import Document

from transcript_task.docx_writer import _truncated, write_docx
from transcript_task.settings import Settings

LONG_DESCRIPTION = "Esta é uma descrição bastante longa. " * 10  # > 255 chars


def _base_item(**overrides) -> dict:
    item = {
        "raw_transcript": "texto bruto",
        "refined_transcript": "texto revisto",
        "duration_seconds": 42.0,
        "source_format": "aac",
        "source_sample_rate": 48000,
        "source_channels": 1,
    }
    item.update(overrides)
    return item


class TestTruncated:
    def test_short_text_is_untouched(self):
        assert _truncated("hello") == "hello"

    def test_long_text_is_capped_at_255(self):
        long_text = "x" * 500
        result = _truncated(long_text)
        assert len(result) <= 255

    def test_boundary_is_not_off_by_one(self):
        assert len(_truncated("x" * 255)) == 255
        assert len(_truncated("x" * 256)) <= 255

    def test_cuts_on_a_word_boundary_not_mid_word(self):
        text = "palavra " * 40  # plenty of spaces near the 255-char cut point
        result = _truncated(text)
        assert result.endswith("…")
        # the character right before the ellipsis must end a whole word --
        # i.e. it's preceded by the start-of-string or a space
        before_ellipsis = result[:-1]
        assert before_ellipsis == "" or text.startswith(before_ellipsis)
        assert not before_ellipsis or text[len(before_ellipsis)] in (" ", "")


class TestWriteDocxWithSummary:
    def test_long_description_does_not_raise(self, tmp_path):
        assert len(LONG_DESCRIPTION) > 255
        item = _base_item(summary={
            "title": "Título de teste",
            "description": LONG_DESCRIPTION,
            "topics": ["a", "b", "c"],
            "speakers_detected": 2,
            "language_variant": "pt-PT",
            "sensitivity": "low",
            "confidence": "high",
        })
        out = tmp_path / "out.docx"

        write_docx("recording.m4a", item, out, Settings())

        assert out.exists()
        doc = Document(str(out))
        assert len(doc.core_properties.subject) <= 255
        # the full, untruncated description must still be in the document body
        assert LONG_DESCRIPTION.strip() in "\n".join(p.text for p in doc.paragraphs)

    def test_sensitivity_banner_present_for_high_sensitivity(self, tmp_path):
        item = _base_item(summary={
            "title": "Assunto sensível",
            "description": "Descrição curta.",
            "topics": ["privado"],
            "speakers_detected": 1,
            "language_variant": "pt-PT",
            "sensitivity": "high",
            "confidence": "high",
        })
        out = tmp_path / "out.docx"
        write_docx("recording.m4a", item, out, Settings())
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "sensível" in text.lower() or "⚠" in text

    def test_no_banner_for_low_sensitivity(self, tmp_path):
        item = _base_item(summary={
            "title": "Assunto normal",
            "description": "Descrição curta.",
            "topics": ["rotina"],
            "speakers_detected": 1,
            "language_variant": "pt-PT",
            "sensitivity": "low",
            "confidence": "high",
        })
        out = tmp_path / "out.docx"
        write_docx("recording.m4a", item, out, Settings())
        text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "⚠" not in text


class TestMetadataAnonymization:
    """Phase 3b: the PII safety net must redact the docx *metadata*
    (core properties) but never the visible body -- a title/description
    with a name in it is exactly as informative to a reviewer who already
    has the document open as one without, since the full transcript with
    the real name is right there a few paragraphs down."""

    def _item_with_name_in_summary(self, **summary_overrides) -> dict:
        summary = {
            "title": "Chamada com Marta sobre dinheiro",
            "description": "Marta discute uma despesa recente com o interlocutor.",
            "topics": ["Marta", "dinheiro"],
            "speakers_detected": 2,
            "language_variant": "pt-PT",
            "sensitivity": "medium",
            "confidence": "high",
        }
        summary.update(summary_overrides)
        return _base_item(summary=summary)

    def test_metadata_is_redacted_by_default(self, tmp_path):
        item = self._item_with_name_in_summary()
        out = tmp_path / "out.docx"
        write_docx("recording.m4a", item, out, Settings())
        doc = Document(str(out))
        assert "marta" not in (doc.core_properties.title or "").lower()
        assert "marta" not in (doc.core_properties.subject or "").lower()
        assert "marta" not in (doc.core_properties.keywords or "").lower()

    def test_visible_body_keeps_the_real_name(self, tmp_path):
        """The point of the exercise: metadata is scrubbed, but the document
        a reviewer actually opens is untouched."""
        item = self._item_with_name_in_summary()
        out = tmp_path / "out.docx"
        write_docx("recording.m4a", item, out, Settings())
        body_text = "\n".join(p.text for p in Document(str(out)).paragraphs)
        assert "Marta" in body_text

    def test_anonymize_metadata_false_keeps_the_real_name_in_metadata_too(self, tmp_path):
        item = self._item_with_name_in_summary()
        out = tmp_path / "out.docx"
        write_docx("recording.m4a", item, out, Settings(anonymize_metadata=False))
        doc = Document(str(out))
        assert "marta" in (doc.core_properties.title or "").lower()


class TestWriteDocxWithoutSummary:
    def test_still_works_without_a_summary(self, tmp_path):
        """Backward compatibility: items produced before Phase 3 (or where
        summarize failed and was never retried) have no 'summary' key."""
        item = _base_item()
        out = tmp_path / "out.docx"

        write_docx("recording.m4a", item, out, Settings())

        assert out.exists()
        doc = Document(str(out))
        assert doc.core_properties.title == "Transcrição — recording.m4a"
