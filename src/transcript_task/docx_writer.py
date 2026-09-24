"""Word document generation for one transcribed recording."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from .settings import Settings


def human_duration(seconds: float) -> str:
    m, s = divmod(int(round(seconds)), 60)
    return f"{m:d}:{s:02d}"


def write_docx(key: str, item: dict, out_path: Path, settings: Settings) -> None:
    """Write one reviewed transcript as a .docx.

    `item` is the corresponding entry from the transcripts.json state dict --
    see pipeline.py's stage_docx for the fields it's expected to carry.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    body = item.get("refined_transcript") or item.get("raw_transcript")
    if not body:
        raise ValueError(f"{key}: no transcript to write")

    doc = Document()
    doc.core_properties.title = f"Transcrição — {key}"
    doc.core_properties.comments = f"Ficheiro de origem: {key}"

    doc.add_heading("Transcrição de Áudio", level=0)

    # Provenance block: this is what ties the document back to its recording.
    meta = doc.add_paragraph()
    meta.add_run("Ficheiro de origem: ").bold = True
    meta.add_run(key)
    meta.add_run("\nDuração: ").bold = True
    meta.add_run(human_duration(item["duration_seconds"])
                 if item.get("duration_seconds") else "desconhecida")
    meta.add_run("\nFormato original: ").bold = True
    meta.add_run(f"{item.get('source_format', '?')} · "
                 f"{item.get('source_sample_rate', '?')} Hz · "
                 f"{item.get('source_channels', '?')} canal(is)")
    meta.add_run("\nModelo de transcrição: ").bold = True
    meta.add_run(settings.asr_model)
    meta.add_run("\nModelo de revisão: ").bold = True
    meta.add_run(settings.llm_model if item.get("refined_transcript")
                 else "— (texto bruto, sem revisão)")
    meta.add_run("\nGerado em: ").bold = True
    meta.add_run(datetime.now().strftime("%Y-%m-%d %H:%M"))
    for run in meta.runs:
        run.font.size = Pt(9)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "Documento gerado automaticamente. Trechos marcados com [?] são "
        "incertos e requerem confirmação humana."
    )
    note_run.italic = True
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    doc.add_paragraph()
    doc.add_heading("Transcrição revista" if item.get("refined_transcript")
                    else "Transcrição (bruta)", level=1)

    for block in [b.strip() for b in body.split("\n") if b.strip()]:
        para = doc.add_paragraph(block)
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.space_after = Pt(8)

    # Keep the unedited ASR output alongside it so a reviewer can check any
    # correction the model made against what was actually heard.
    if item.get("refined_transcript") and item.get("raw_transcript"):
        doc.add_page_break()
        doc.add_heading("Anexo — transcrição automática original", level=1)
        anexo = doc.add_paragraph(item["raw_transcript"])
        anexo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in anexo.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.save(out_path)
